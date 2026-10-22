"""Builds the Private AI 3-year monthly financial model (xlsx).

Run: python3 build_model.py [--calib calib.json]
All projections are Excel formulas; inputs live on the Inputs and Schedules sheets.
"""
import json
import sys
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

OUT = "private-ai-3yr-model-2026-10-20.xlsx"
SCEN = ["Original plan", "Conservative", "Base", "Stretch"]
SHEETS = ["M_Original", "M_Conservative", "M_Base", "M_Stretch"]
N = 36
C0 = 3  # first month column (C)

calib = {"unitemized": [0, 0, 0], "orig_projects": [0, 0, 0]}
if "--calib" in sys.argv:
    calib = json.load(open(sys.argv[sys.argv.index("--calib") + 1]))

F = "Arial"
BLUE = Font(name=F, color="0000FF")
BLACK = Font(name=F, color="000000")
GREEN = Font(name=F, color="008000")
BOLD = Font(name=F, bold=True)
HDR = Font(name=F, bold=True, color="FFFFFF")
HFILL = PatternFill("solid", fgColor="1F3864")
YELLOW = PatternFill("solid", fgColor="FFFF00")
SEC = PatternFill("solid", fgColor="D9E1F2")
USD = '$#,##0;($#,##0);"-"'
PCT = '0.0%;(0.0%);"-"'
NUM = '#,##0.0;(#,##0.0);"-"'
INT = '#,##0;(#,##0);"-"'
thin = Side(style="thin", color="BFBFBF")

wb = Workbook()

# ---------------------------------------------------------------- Inputs
inp = wb.active
inp.title = "Inputs"
inp["A1"] = "Inputs — Private AI 3-year model (as of 2026-10-20)"
inp["A1"].font = Font(name=F, bold=True, size=14)
inp["A2"] = ("Blue = input you can change. Yellow fill = key assumption or a value still to confirm. "
             "Column G gives the source or basis for every number.")
inp["A2"].font = Font(name=F, italic=True)
heads = ["Parameter", "Unit"] + SCEN + ["Source / basis"]
for i, h in enumerate(heads, 1):
    c = inp.cell(row=4, column=i, value=h)
    c.font, c.fill = HDR, HFILL
SRC = "Founder brief, prompts/init.md"
R14 = "Analysis, 14-revised-plan-2026-10-20.md"

# key, label, unit, [orig, cons, base, stretch], fmt, source, key_assumption
P = [
    ("sec", "Professional-firm clients"),
    ("twk_sign", "TWK signing month", "month", [1, 1, 1, 1], INT, "Assumption: TWK pilot starts week 1 (Oct 19, 2026)", False),
    ("twk_setup", "TWK setup fee (after discount)", "$", [20000, 20000, 20000, 20000], USD, "Assumption; discount amount not in source (TODO)", True),
    ("twk_sub", "TWK subscription", "$/month", [3800, 3000, 3800, 3800], USD, "Assumption; TWK tier not in source (TODO)", True),
    ("setup", "Setup fee per other client", "$", [25000, 15000, 19000, 25000], USD, f"{SRC}: unit economics $25,000, range $15,000–40,000. Base uses the original model's implied ~$19,100; conservative the range floor; stretch the unit-economics figure", True),
    ("sub", "Subscription per other client (blended)", "$/month", [4500, 2500, 2500, 3500], USD, f"{SRC}: Starter $2,500, Professional $4,500–6,500. Base uses the original model's implied ~$2,490 average; stretch a half-Starter, half-Professional blend", True),
    ("pct_sign", "Share of setup billed at signing", "%", [0, 0.5, 0.5, 0.5], PCT, "Decision Oct 20, 2026 (50% at signing); original assumed billed at install", False),
    ("lag", "Months from signing to acceptance", "months", [2, 2, 2, 2], INT, "Assumption; Phase 1 runs 12 weeks", False),
    ("esc", "Annual subscription escalator", "%", [0.05, 0.05, 0.05, 0.05], PCT, f"{SRC}: 5% annual escalator (applied per model year)", False),
    ("addon", "Add-on revenue per adopting client", "$/month", [1000, 1000, 1000, 1000], USD, f"{SRC}: $12,000/yr per client in Years 2–3", False),
    ("adopt", "Share of clients buying add-ons", "%", [1, 0.25, 0.5, 0.75], PCT, f"Original: unit economics assume 1–2 add-ons; revised shares are {R14}", True),
    ("addon_lag", "Add-ons start (months after acceptance)", "months", [3, 3, 3, 3], INT, f"{SRC}: add-ons after the first quarter", False),
    ("hw_cost", "Hardware cost per install", "$", [28000, 28000, 28000, 28000], USD, f"{SRC}: unit economics $28,000", False),
    ("hw_pct", "Hardware pass-through margin", "%", [0.08, 0.08, 0.08, 0.08], PCT, f"{SRC}: cost plus 8%", False),
    ("sec", "Industrial"),
    ("comp_setup", "Contractor compliance setup", "$", [0, 3000, 4000, 5000], USD, f"{SRC}: $2,500–6,000; case values per {R14}", False),
    ("comp_ret", "Contractor compliance retainer", "$/month", [0, 400, 500, 600], USD, f"{SRC}: $400–900/month; expect low end until proven ({R14})", True),
    ("sec", "Gross margin by line"),
    ("gm_setup", "Setup", "%", [0.55] * 4, PCT, f"{SRC}", False),
    ("gm_sub", "Subscription", "%", [0.70] * 4, PCT, f"{SRC}", False),
    ("gm_add", "Workflow add-ons", "%", [0.85] * 4, PCT, f"{SRC}", False),
    ("gm_hw", "Hardware margin", "%", [1.0] * 4, PCT, f"{SRC}", False),
    ("gm_comp", "Contractor compliance", "%", [0.80] * 4, PCT, f"{SRC}: 75–85% (midpoint)", False),
    ("gm_proj", "Industrial projects", "%", [0.65] * 4, PCT, f"{SRC}: services range 55–85%; weekly target 60%+ (assumed 65%)", True),
    ("sec", "People"),
    ("sal1", "Founder salary, Year 1", "$/yr", [70000] * 4, USD, f"{SRC}", False),
    ("sal2", "Founder salary, Year 2", "$/yr", [100000] * 4, USD, f"{SRC}", False),
    ("sal3", "Founder salary, Year 3", "$/yr", [130000] * 4, USD, f"{SRC}", False),
    ("fica", "Employer payroll taxes on founder salary", "%", [0.0765] * 4, PCT, "Employer FICA rate (Social Security 6.2% + Medicare 1.45%)", False),
    ("hire1", "First hire: start month", "month", [10, 26, 17, 14], INT, f"Original: months 9–12 ({SRC}). Revised: the month after MRR first reaches ~$20K in that scenario (see Summary), per {R14}", True),
    ("cost1", "First hire: loaded annual cost", "$/yr", [82500, 87500, 87500, 87500], USD, f"Original ~$70–95K ({SRC}); revised ~$74–101K at 1.35x (BLS ECEC) — midpoints", False),
    ("hire2", "Second hire: start month", "month", [24, 37, 29, 22], INT, f"Original months 22–26 ({SRC}). Revised: when MRR reaches ~$45K in that scenario (37 = not within the model horizon), per {R14}", True),
    ("cost2", "Second hire: loaded annual cost", "$/yr", [97500, 105000, 105000, 105000], USD, f"Original ~$80–115K ({SRC}); revised $65–90K base x1.35 — midpoints", False),
    ("ctr_start", "Contract engineer + student: start month", "month", [37, 6, 6, 6], INT, f"{R14}; ends when the first hire starts; original plan had none", False),
    ("ctr", "Contract engineer", "$/month", [0, 1400, 1400, 1400], USD, f"{R14}: $600–2,200/month (midpoint)", False),
    ("stu", "Student worker", "$/month", [0, 900, 900, 900], USD, f"{R14}: $15–20/hr, 10–15 hrs/week (approx. midpoint)", False),
    ("sec", "Other operating costs"),
    ("mkt", "Memberships, events, travel, collateral", "$/yr", [25000] * 4, USD, f"{SRC}: about $25K a year", False),
    ("ins", "Business insurance", "$/yr", [8000] * 4, USD, f"{SRC}: $6–10K in Year 1 (midpoint)", False),
    ("keyins", "Key-person life and disability cover", "$/yr", [0, 1200, 1200, 1200], USD, f"Assumption; recommended in {R14}", True),
    ("sw", "Software, tools, hosting, backups", "$/month", [1000] * 4, USD, "Assumption (not in source)", True),
    ("prof", "Bookkeeping, legal, accounting", "$/month", [1000] * 4, USD, "Assumption (not in source)", True),
    ("formation", "One-time legal and formation (months 1–3)", "$", [12000] * 4, USD, f"{SRC}: $20K use-of-funds line for insurance, legal, formation, SDVOSB, less insurance", False),
    ("un1", "Unitemized costs, Year 1 (calibrated)", "$/yr", [calib["unitemized"][0]] * 4, USD, "Calibrated so the Original plan scenario matches the source operating results; itemize when the original spreadsheet is found", True),
    ("un2", "Unitemized costs, Year 2 (calibrated)", "$/yr", [calib["unitemized"][1]] * 4, USD, "As above", True),
    ("un3", "Unitemized costs, Year 3 (calibrated)", "$/yr", [calib["unitemized"][2]] * 4, USD, "As above", True),
    ("sec", "Capex and funding"),
    ("demo_m", "Demo/lab appliance: purchase month", "month", [1] * 4, INT, "Assumption", False),
    ("demo", "Demo/lab appliance, tools, software", "$", [35000] * 4, USD, f"{SRC}: use of funds $35,000", False),
    ("fcap", "Founder capital", "$", [0] * 4, USD, "TODO: amount not in source", True),
    ("debt_m", "Debt draw month", "month", [5] * 4, INT, f"{SRC}: debt and equity close in months 3–6", False),
    ("debt", "Debt amount", "$", [200000] * 4, USD, f"{SRC}: $150–250K (midpoint)", False),
    ("rate", "Debt interest rate", "%", [0.105] * 4, PCT, "Assumption; SBA 7(a) terms to verify", True),
    ("term", "Debt term", "months", [120] * 4, INT, "Assumption (10-year term)", True),
    ("eq_m", "Equity close month", "month", [5] * 4, INT, f"{SRC}: months 3–6", False),
    ("eq", "Equity amount", "$", [200000] * 4, USD, f"{SRC}: $150–250K (midpoint)", False),
    ("buffer_m", "Operating buffer (months of Year-1 average costs)", "months", [6] * 4, INT, f"{SRC}: six-month operating buffer", False),
]
ORIG = calib.get("orig", {})
P = [(p[0], p[1], p[2], [ORIG.get(p[0], p[3][0])] + p[3][1:], p[4],
      (p[5] + f" | Original column calibrated to the brief's Year-3 revenue mix (Checks)" if p[0] in ORIG else p[5]),
      p[6] or p[0] in ORIG)
     if p[0] != "sec" else p for p in P]
ROW = {}
r = 5
for p in P:
    if p[0] == "sec":
        c = inp.cell(row=r, column=1, value=p[1])
        c.font = BOLD
        for col in range(1, 8):
            inp.cell(row=r, column=col).fill = SEC
        r += 1
        continue
    key, label, unit, vals, fmt, src, keyass = p
    ROW[key] = r
    inp.cell(row=r, column=1, value=label).font = BLACK
    inp.cell(row=r, column=2, value=unit).font = BLACK
    for j, v in enumerate(vals):
        c = inp.cell(row=r, column=3 + j, value=v)
        c.font, c.number_format = BLUE, fmt
        if keyass:
            c.fill = YELLOW
    inp.cell(row=r, column=7, value=src).font = BLACK
    r += 1
# derived: monthly debt payment
ROW["pmt"] = r
inp.cell(row=r, column=1, value="Monthly debt payment (formula)").font = BLACK
inp.cell(row=r, column=2, value="$/month")
for j in range(4):
    col = L(3 + j)
    c = inp.cell(row=r, column=3 + j,
                 value=f"=IF({col}{ROW['debt']}=0,0,PMT({col}{ROW['rate']}/12,{col}{ROW['term']},-{col}{ROW['debt']}))")
    c.font, c.number_format = BLACK, USD
inp.cell(row=r, column=7, value="Standard amortizing payment")
inp.column_dimensions["A"].width = 46
inp.column_dimensions["B"].width = 10
for col in "CDEF":
    inp.column_dimensions[col].width = 15
inp.column_dimensions["G"].width = 95
inp.freeze_panes = "C5"

# ------------------------------------------------------------- Schedules
sch = wb.create_sheet("Schedules")
sch["A1"] = "Schedules — monthly signings and project revenue by scenario (inputs)"
sch["A1"].font = Font(name=F, bold=True, size=14)
sch["A2"] = ("Month 1 starts Oct 19, 2026. CPA firms close May–December (tax season Jan–Apr); law firms year-round. "
             "Professional signings exclude TWK. Edit the blue cells.")
sch["A2"].font = Font(name=F, italic=True)
sch.cell(row=4, column=1, value="Row").font = HDR
sch.cell(row=4, column=2, value="Total").font = HDR
for m in range(1, N + 1):
    c = sch.cell(row=4, column=C0 + m - 1, value=m)
    c.font, c.fill = HDR, HFILL
sch["A4"].fill = sch["B4"].fill = HFILL


def months(lst):
    return {m: 1 for m in lst}


def spread(total, ms):
    return {m: total / len(ms) for m in ms}


op = calib["orig_projects"]
schedule = {
    "Original plan": {
        "prof": months([3, 6, 9, 13, 15, 17, 19, 20, 22, 23, 25, 26, 28, 29, 30, 32, 33, 34, 35]),
        "comp": {},
        "proj": {**spread(op[0], range(1, 13)), **spread(op[1], range(13, 25)), **spread(op[2], range(25, 37))},
        "note": "Original: 4 clients in Year 1 incl. TWK, 7 in Year 2, 9 in Year 3, ignoring the CPA calendar; industrial project revenue calibrated to the source revenue totals",
    },
    "Conservative": {
        "prof": months([7, 10, 13, 17, 20, 22, 24, 25, 28, 31, 33, 34, 36]),
        "comp": months([3, 6, 9, 11, 14, 16, 18, 20, 22, 26, 30, 34]),
        "proj": {20: 15000, 30: 20000},
        "note": "Conservative per 14-revised-plan",
    },
    "Base": {
        "prof": months([6, 8, 10, 13, 14, 16, 18, 20, 21, 23, 25, 26, 28, 30, 31, 32, 33, 34, 35]),
        "comp": months([2, 3, 5, 7, 9, 11, 13, 14, 15, 16, 17, 19, 21, 23, 25, 27, 29, 31]),
        "proj": {6: 5000, 8: 4000, 19: 20000, 28: 20000, 33: 20000},
        "note": "Base per 14-revised-plan: 3 beyond TWK in Year 1, then 7 and 9 (plan counts)",
    },
    "Stretch": {
        "prof": months([5, 8, 9, 10, 13, 14, 15, 17, 19, 20, 21, 22, 23, 25, 26, 27, 28, 29, 31, 32, 33, 34, 35, 36]),
        "comp": months([2, 3, 4, 5, 7, 8, 9, 11] + list(range(13, 23)) + [25, 27, 29, 31, 33, 35]),
        "proj": {6: 5000, 8: 4000, 10: 20000, 16: 25000, 19: 20000, 22: 30000, 27: 30000, 30: 40000, 33: 30000},
        "note": "Stretch per 14-revised-plan",
    },
}
SROW = {}
r = 5
for s in SCEN:
    c = sch.cell(row=r, column=1, value=s)
    c.font = BOLD
    for col in range(1, C0 + N):
        sch.cell(row=r, column=col).fill = SEC
    sch.cell(row=r, column=C0, value=schedule[s]["note"]).font = Font(name=F, italic=True)
    r += 1
    for kind, label, fmt in [("prof", "Professional clients signed (excl. TWK)", INT),
                             ("comp", "Contractor compliance clients signed", INT),
                             ("proj", "Industrial project revenue ($)", USD)]:
        SROW[(s, kind)] = r
        sch.cell(row=r, column=1, value=label).font = BLACK
        tot = sch.cell(row=r, column=2, value=f"=SUM({L(C0)}{r}:{L(C0 + N - 1)}{r})")
        tot.font, tot.number_format = BLACK, fmt
        for m in range(1, N + 1):
            c = sch.cell(row=r, column=C0 + m - 1, value=schedule[s][kind].get(m, 0))
            c.font, c.number_format = BLUE, fmt
        r += 1
    r += 1
sch.column_dimensions["A"].width = 40
sch.column_dimensions["B"].width = 11
for m in range(N):
    sch.column_dimensions[L(C0 + m)].width = 9
sch.freeze_panes = "C5"

# --------------------------------------------------------- Monthly sheets
LINES = [
    ("hdr", "Clients"),
    ("twk_signed", "TWK signed", INT),
    ("signed", "Professional clients signed (excl. TWK)", INT),
    ("accepted", "Professional clients accepted (excl. TWK)", INT),
    ("twk_acc", "TWK accepted", INT),
    ("cum_acc", "Cumulative accepted (incl. TWK)", INT),
    ("active", "Subscribing clients (excl. TWK)", NUM),
    ("twk_active", "TWK subscribing", INT),
    ("addon_cl", "Clients eligible for add-ons", INT),
    ("comp_new", "Compliance clients signed", INT),
    ("comp_act", "Compliance clients on retainer", INT),
    ("hdr", "Revenue ($)"),
    ("escf", "Escalator factor", '0.000'),
    ("rev_sub", "Subscriptions", USD),
    ("rev_add", "Workflow add-ons", USD),
    ("rev_setup", "Setup and implementation", USD),
    ("rev_hw", "Hardware margin", USD),
    ("rev_comp", "Contractor compliance", USD),
    ("rev_proj", "Industrial projects", USD),
    ("rev", "Total revenue", USD),
    ("mrr", "MRR (subscriptions + add-ons + retainers)", USD),
    ("hdr", "Gross profit ($)"),
    ("cogs", "Cost of revenue", USD),
    ("gp", "Gross profit", USD),
    ("hdr", "Operating costs ($)"),
    ("o_sal", "Founder salary", USD),
    ("o_tax", "Employer payroll taxes (founder)", USD),
    ("o_staff", "Staff (loaded)", USD),
    ("o_ctr", "Contract engineer + student", USD),
    ("o_mkt", "Memberships, events, travel", USD),
    ("o_ins", "Insurance (incl. key-person)", USD),
    ("o_sw", "Software, tools, hosting", USD),
    ("o_prof", "Professional fees + formation", USD),
    ("o_un", "Unitemized (calibrated)", USD),
    ("opex", "Total operating costs", USD),
    ("opres", "Operating result", USD),
    ("ttm", "Operating result, trailing 12 months", USD),
    ("hdr", "Cash ($)"),
    ("capex", "Capex (demo appliance)", USD),
    ("f_cap", "Founder capital in", USD),
    ("f_debt", "Debt drawn", USD),
    ("f_eq", "Equity in", USD),
    ("d_open", "Debt balance, opening", USD),
    ("d_int", "Interest", USD),
    ("d_pmt", "Debt payment", USD),
    ("d_close", "Debt balance, closing", USD),
    ("ncf", "Net cash flow", USD),
    ("cash", "Cash balance, end of month", USD),
    ("cum_pre", "Cumulative cash before funding (after capex and debt service)", USD),
]


def build_month_sheet(name, si):
    ws = wb.create_sheet(name)
    s = SCEN[si]
    icol = L(3 + si)  # scenario column on Inputs

    def I(key):  # absolute link to Inputs
        return f"Inputs!${icol}${ROW[key]}"

    ws["A1"] = f"{s} — monthly model (all cells are formulas)"
    ws["A1"].font = Font(name=F, bold=True, size=14)
    ws.cell(row=2, column=1, value="Month").font = BOLD
    ws.cell(row=3, column=1, value="Model year").font = BOLD
    ws.cell(row=4, column=1, value="Month starts").font = BOLD
    for m in range(1, N + 1):
        col = C0 + m - 1
        ws.cell(row=2, column=col, value=m).font = BLUE
        ws.cell(row=3, column=col, value=f"=ROUNDUP({L(col)}2/12,0)").font = BLACK
        d = ws.cell(row=4, column=col, value=f"=DATE(2026,10,19)+ROUND(({L(col)}2-1)*30.44,0)")
        d.font, d.number_format = BLACK, "mmm yy"
        ws.column_dimensions[L(col)].width = 11
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 4
    R = {}
    r = 6
    for ln in LINES:
        if ln[0] == "hdr":
            c = ws.cell(row=r, column=1, value=ln[1])
            c.font = BOLD
            for col in range(1, C0 + N):
                ws.cell(row=r, column=col).fill = SEC
            r += 1
            continue
        R[ln[0]] = r
        r += 1
    first, last = L(C0), L(C0 + N - 1)

    def rowrng(k):
        return f"${first}${R[k]}:${last}${R[k]}"

    for ln in LINES:
        if ln[0] == "hdr":
            continue
        k, label, fmt = ln
        ws.cell(row=R[k], column=1, value=label).font = BOLD if k in ("rev", "gp", "opex", "opres", "mrr", "cash") else BLACK
        for m in range(1, N + 1):
            col = C0 + m - 1
            X, P_ = L(col), L(col - 1)
            mm = f"{X}$2"
            yr = f"{X}$3"
            prev = (lambda kk: f"{P_}{R[kk]}") if m > 1 else (lambda kk: "0")
            green = False
            if k == "twk_signed":
                f = f"=IF({mm}={I('twk_sign')},1,0)"
            elif k == "signed":
                f = f"=Schedules!{X}{SROW[(s, 'prof')]}"; green = True
            elif k == "accepted":
                f = f"=IF({mm}-{I('lag')}>=1,INDEX({rowrng('signed')},{mm}-{I('lag')}),0)"
            elif k == "twk_acc":
                f = f"=IF({mm}={I('twk_sign')}+{I('lag')},1,0)"
            elif k == "cum_acc":
                f = f"={prev('cum_acc')}+{X}{R['accepted']}+{X}{R['twk_acc']}"
            elif k == "active":
                f = f"={prev('active')}+{prev('accepted')}"
            elif k == "twk_active":
                f = f"=IF({mm}>{I('twk_sign')}+{I('lag')},1,0)"
            elif k == "addon_cl":
                f = f"=IF({mm}-{I('addon_lag')}>=1,INDEX({rowrng('cum_acc')},{mm}-{I('addon_lag')}),0)"
            elif k == "comp_new":
                f = f"=Schedules!{X}{SROW[(s, 'comp')]}"; green = True
            elif k == "comp_act":
                f = f"={prev('comp_act')}+{prev('comp_new')}"
            elif k == "escf":
                f = f"=(1+{I('esc')})^({yr}-1)"
            elif k == "rev_sub":
                f = f"=({X}{R['active']}*{I('sub')}+{X}{R['twk_active']}*{I('twk_sub')})*{X}{R['escf']}"
            elif k == "rev_add":
                f = f"={X}{R['addon_cl']}*{I('addon')}*{I('adopt')}*{X}{R['escf']}"
            elif k == "rev_setup":
                f = (f"={X}{R['signed']}*{I('setup')}*{I('pct_sign')}+{X}{R['accepted']}*{I('setup')}*(1-{I('pct_sign')})"
                     f"+{X}{R['twk_signed']}*{I('twk_setup')}*{I('pct_sign')}+{X}{R['twk_acc']}*{I('twk_setup')}*(1-{I('pct_sign')})")
            elif k == "rev_hw":
                f = f"=({X}{R['accepted']}+{X}{R['twk_acc']})*{I('hw_cost')}*{I('hw_pct')}"
            elif k == "rev_comp":
                f = f"={X}{R['comp_new']}*{I('comp_setup')}+{X}{R['comp_act']}*{I('comp_ret')}"
            elif k == "rev_proj":
                f = f"=Schedules!{X}{SROW[(s, 'proj')]}"; green = True
            elif k == "rev":
                f = f"=SUM({X}{R['rev_sub']}:{X}{R['rev_proj']})"
            elif k == "mrr":
                f = f"={X}{R['rev_sub']}+{X}{R['rev_add']}+{X}{R['comp_act']}*{I('comp_ret')}"
            elif k == "cogs":
                f = (f"={X}{R['rev_setup']}*(1-{I('gm_setup')})+{X}{R['rev_sub']}*(1-{I('gm_sub')})"
                     f"+{X}{R['rev_add']}*(1-{I('gm_add')})+{X}{R['rev_hw']}*(1-{I('gm_hw')})"
                     f"+{X}{R['rev_comp']}*(1-{I('gm_comp')})+{X}{R['rev_proj']}*(1-{I('gm_proj')})")
            elif k == "gp":
                f = f"={X}{R['rev']}-{X}{R['cogs']}"
            elif k == "o_sal":
                f = f"=CHOOSE({yr},{I('sal1')},{I('sal2')},{I('sal3')})/12"
            elif k == "o_tax":
                f = f"={X}{R['o_sal']}*{I('fica')}"
            elif k == "o_staff":
                f = f"=IF({mm}>={I('hire1')},{I('cost1')}/12,0)+IF({mm}>={I('hire2')},{I('cost2')}/12,0)"
            elif k == "o_ctr":
                f = f"=IF(AND({mm}>={I('ctr_start')},{mm}<{I('hire1')}),{I('ctr')}+{I('stu')},0)"
            elif k == "o_mkt":
                f = f"={I('mkt')}/12"
            elif k == "o_ins":
                f = f"=({I('ins')}+{I('keyins')})/12"
            elif k == "o_sw":
                f = f"={I('sw')}"
            elif k == "o_prof":
                f = f"={I('prof')}+IF({mm}<=3,{I('formation')}/3,0)"
            elif k == "o_un":
                f = f"=CHOOSE({yr},{I('un1')},{I('un2')},{I('un3')})/12"
            elif k == "opex":
                f = f"=SUM({X}{R['o_sal']}:{X}{R['o_un']})"
            elif k == "opres":
                f = f"={X}{R['gp']}-{X}{R['opex']}"
            elif k == "ttm":
                f = f"=SUM(INDEX({rowrng('opres')},MAX(1,{mm}-11)):{X}{R['opres']})"
            elif k == "capex":
                f = f"=IF({mm}={I('demo_m')},{I('demo')},0)"
            elif k == "f_cap":
                f = f"=IF({mm}=1,{I('fcap')},0)"
            elif k == "f_debt":
                f = f"=IF({mm}={I('debt_m')},{I('debt')},0)"
            elif k == "f_eq":
                f = f"=IF({mm}={I('eq_m')},{I('eq')},0)"
            elif k == "d_open":
                f = f"={prev('d_close')}"
            elif k == "d_int":
                f = f"={X}{R['d_open']}*{I('rate')}/12"
            elif k == "d_pmt":
                f = f"=IF({mm}>{I('debt_m')},MIN(Inputs!${icol}${ROW['pmt']},{X}{R['d_open']}+{X}{R['d_int']}),0)"
            elif k == "d_close":
                f = f"={X}{R['d_open']}+{X}{R['f_debt']}+{X}{R['d_int']}-{X}{R['d_pmt']}"
            elif k == "ncf":
                f = (f"={X}{R['opres']}-{X}{R['capex']}+{X}{R['f_cap']}+{X}{R['f_debt']}+{X}{R['f_eq']}"
                     f"-{X}{R['d_pmt']}")
            elif k == "cash":
                f = f"={prev('cash')}+{X}{R['ncf']}"
            elif k == "cum_pre":
                f = f"={prev('cum_pre')}+{X}{R['opres']}-{X}{R['capex']}-{X}{R['d_pmt']}"
            c = ws.cell(row=R[k], column=col, value=f)
            c.font = GREEN if green else BLACK
            c.number_format = fmt
    ws.freeze_panes = "C5"
    return R


MROWS = {}
for i, name in enumerate(SHEETS):
    MROWS[name] = build_month_sheet(name, i)

# ---------------------------------------------------------------- Summary
sm = wb.create_sheet("Summary", 0)
sm["A1"] = "Private AI — 3-year model summary (as of 2026-10-20)"
sm["A1"].font = Font(name=F, bold=True, size=14)
sm["A2"] = ("Revised scenarios follow 14-revised-plan-2026-10-20.md; 'Original plan' reproduces the founder brief's figures "
            "(see Checks). All values are formulas; change inputs on the Inputs and Schedules sheets.")
sm["A2"].font = Font(name=F, italic=True)
sm.cell(row=4, column=1, value="Metric").font = HDR
sm.cell(row=4, column=1).fill = HFILL
for j, s in enumerate(SCEN):
    c = sm.cell(row=4, column=2 + j, value=s)
    c.font, c.fill = HDR, HFILL
R0 = MROWS[SHEETS[0]]
fl, ll = L(C0), L(C0 + N - 1)


def yrsum(sheet, key, y):
    rr = MROWS[sheet][key]
    return f"SUMIF({sheet}!${fl}$3:${ll}$3,{y},{sheet}!${fl}${rr}:${ll}${rr})"


def at(sheet, key, m):
    return f"{sheet}!{L(C0 + m - 1)}{MROWS[sheet][key]}"


rows = []
for y in (1, 2, 3):
    rows += [(f"Revenue, Year {y}", lambda sh, y=y: f"={yrsum(sh, 'rev', y)}", USD)]
for y in (1, 2, 3):
    rows += [(f"Gross margin, Year {y}", lambda sh, y=y: f"=IFERROR({yrsum(sh, 'gp', y)}/{yrsum(sh, 'rev', y)},0)", PCT)]
for y in (1, 2, 3):
    rows += [(f"Operating result, Year {y}", lambda sh, y=y: f"={yrsum(sh, 'opres', y)}", USD)]
for y in (1, 2, 3):
    rows += [(f"Industrial revenue, Year {y}",
              lambda sh, y=y: f"={yrsum(sh, 'rev_comp', y)}+{yrsum(sh, 'rev_proj', y)}", USD)]
for m in (12, 24, 36):
    rows += [(f"Professional clients incl. TWK, month {m} (signed)",
              lambda sh, m=m: f"=SUM({sh}!${fl}${MROWS[sh]['signed']}:{L(C0 + m - 1)}{MROWS[sh]['signed']})+SUM({sh}!${fl}${MROWS[sh]['twk_signed']}:{L(C0 + m - 1)}{MROWS[sh]['twk_signed']})", INT)]
for m in (6, 9, 12, 24, 36):
    rows += [(f"MRR, month {m}", lambda sh, m=m: f"={at(sh, 'mrr', m)}", USD)]
rows += [("First month MRR reaches $20,000",
          lambda sh: f'=IFERROR(INDEX({sh}!${fl}$2:${ll}$2,MATCH(TRUE,INDEX({sh}!${fl}${MROWS[sh]["mrr"]}:${ll}${MROWS[sh]["mrr"]}>=20000,0),0)),"after month 36")', INT)]
rows += [("Breakeven: first month the trailing-12-month operating result is positive",
          lambda sh: f'=IFERROR(INDEX({sh}!${fl}$2:${ll}$2,MATCH(TRUE,INDEX({sh}!${fl}${MROWS[sh]["ttm"]}:${ll}${MROWS[sh]["ttm"]}>0,0),0)),"after month 36")', INT)]
rows += [("Lowest cumulative cash before funding (trough)",
          lambda sh: f"=MIN({sh}!${fl}${MROWS[sh]['cum_pre']}:${ll}${MROWS[sh]['cum_pre']})", USD)]
rows += [("Operating buffer (months x Year-1 average monthly costs)",
          lambda sh: f"=Inputs!{L(3 + SHEETS.index(sh))}{ROW['buffer_m']}*{yrsum(sh, 'opex', 1)}/12", USD)]
rows += [("Capital needed (trough + buffer)", None, USD)]
rows += [("Planned raise (debt + equity + founder capital)",
          lambda sh: f"=Inputs!{L(3 + SHEETS.index(sh))}{ROW['debt']}+Inputs!{L(3 + SHEETS.index(sh))}{ROW['eq']}+Inputs!{L(3 + SHEETS.index(sh))}{ROW['fcap']}", USD)]
rows += [("Lowest cash balance with the planned raise (negative = bridge or founder capital needed before funding closes)",
          lambda sh: f"=MIN({sh}!${fl}${MROWS[sh]['cash']}:${ll}${MROWS[sh]['cash']})", USD)]
rows += [("Cash balance, month 36", lambda sh: f"={at(sh, 'cash', 36)}", USD)]

SMROW = {}
r = 5
for label, fn, fmt in rows:
    SMROW[label] = r
    sm.cell(row=r, column=1, value=label).font = BLACK
    for j, sh in enumerate(SHEETS):
        if fn is None:
            colL = L(2 + j)
            f = f"=-{colL}{SMROW['Lowest cumulative cash before funding (trough)']}+{colL}{SMROW['Operating buffer (months x Year-1 average monthly costs)']}"
        else:
            f = fn(sh)
        c = sm.cell(row=r, column=2 + j, value=f)
        c.font, c.number_format = GREEN, fmt
    r += 1
sm.column_dimensions["A"].width = 60
for j in range(4):
    sm.column_dimensions[L(2 + j)].width = 17
sm.freeze_panes = "B5"

# ----------------------------------------------------------------- Checks
ck = wb.create_sheet("Checks", 1)
ck["A1"] = "Checks — Original plan scenario vs the founder brief"
ck["A1"].font = Font(name=F, bold=True, size=14)
ck["A2"] = ("The original spreadsheet isn't available. The Original plan scenario is calibrated (industrial project revenue and "
            "unitemized costs) so revenue and operating results match the brief. Gross margin is a test of the line margins.")
ck["A2"].font = Font(name=F, italic=True)
for i, h in enumerate(["Metric", "Source figure", "Model (Original plan)", "Difference", "Source"], 1):
    c = ck.cell(row=4, column=i, value=h)
    c.font, c.fill = HDR, HFILL
targets = [("Revenue, Year 1", 155000, USD), ("Revenue, Year 2", 465000, USD), ("Revenue, Year 3", 860000, USD),
           ("Gross margin, Year 1", 0.66, PCT), ("Gross margin, Year 2", 0.68, PCT), ("Gross margin, Year 3", 0.69, PCT),
           ("Operating result, Year 1", -95000, USD), ("Operating result, Year 2", 15000, USD),
           ("Operating result, Year 3", 155000, USD), ("Industrial revenue, Year 3", 129000, USD)]
for i, (label, tv, fmt) in enumerate(targets):
    rr = 5 + i
    ck.cell(row=rr, column=1, value=label).font = BLACK
    c = ck.cell(row=rr, column=2, value=tv); c.font, c.number_format = BLUE, fmt
    c = ck.cell(row=rr, column=3, value=f"=Summary!B{SMROW[label]}"); c.font, c.number_format = GREEN, fmt
    c = ck.cell(row=rr, column=4, value=f"=C{rr}-B{rr}"); c.font, c.number_format = BLACK, fmt
    ck.cell(row=rr, column=5, value="Founder brief, section K" if "Industrial" not in label else "Brief: Year-3 mix 15% of ~$860K").font = BLACK
ck.column_dimensions["A"].width = 32
for col in "BCD":
    ck.column_dimensions[col].width = 20
ck.column_dimensions["E"].width = 36

# ----------------------------------------------------------------- README
rd = wb.create_sheet("README", 0)
lines = [
    ("Private AI — 3-year financial model (regenerated 2026-10-20)", Font(name=F, bold=True, size=14)),
    ("", BLACK),
    ("What this is", BOLD),
    ("A monthly, formula-driven model for 36 months from Oct 19, 2026, with four scenarios side by side.", BLACK),
    ("It replaces the original 3-year spreadsheet, which wasn't available; the 'Original plan' scenario is calibrated to that model's published results.", BLACK),
    ("", BLACK),
    ("Sheets", BOLD),
    ("Summary — key results for all four scenarios.", BLACK),
    ("Checks — Original plan scenario vs the figures in the founder brief.", BLACK),
    ("Inputs — every assumption, per scenario, with its source in column G.", BLACK),
    ("Schedules — month-by-month signings and industrial project revenue per scenario.", BLACK),
    ("M_Original / M_Conservative / M_Base / M_Stretch — the monthly models (formulas only).", BLACK),
    ("", BLACK),
    ("Colour code", BOLD),
    ("Blue text = input. Black = formula. Green = link to another sheet. Yellow fill = key assumption or value to confirm.", BLACK),
    ("", BLACK),
    ("How to use", BOLD),
    ("Change blue cells on Inputs or Schedules only. Every other number recalculates.", BLACK),
    ("", BLACK),
    ("Main simplifications", BOLD),
    ("1. Subscriptions start the month after acceptance; the 5% escalator is applied per model year, not per client anniversary.", BLACK),
    ("2. No client churn. Assessment fees are treated as credited to setup.", BLACK),
    ("3. Operating result excludes interest; debt service and capex appear in cash only.", BLACK),
    ("4. 'Unitemized costs' are calibrated so the Original plan scenario matches the source operating results; they apply to every scenario until the original cost lines are found.", BLACK),
    ("5. Capital needed = lowest cumulative cash before funding (after capex and debt service) plus a six-month buffer of Year-1 average costs. The source's own $375K decomposition doesn't reconcile (see 12-open-questions.md, C6).", BLACK),
    ("", BLACK),
    ("Sources", BOLD),
    ("Founder brief: private-ai/prompts/init.md. Revisions and research: private-ai/14-revised-plan-2026-10-20.md.", BLACK),
    ("Built by private-ai/model/build_model.py.", BLACK),
]
for i, (t, fnt) in enumerate(lines, 1):
    c = rd.cell(row=i, column=1, value=t)
    c.font = fnt
rd.column_dimensions["A"].width = 150

for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if c.font and c.font.name != F:
                c.font = Font(name=F, bold=c.font.bold, italic=c.font.italic, color=c.font.color, size=c.font.size)

wb.save(OUT)
print("saved", OUT, "rows:", {k: v for k, v in list(SMROW.items())[:3]})
